"""Penyimpanan data dashboard: import, penggabungan antarperiode, verifikasi, dan backup.

Hasil model (T², UCL, sinyal, tekanan, kontribusi) bersifat read-only: hanya berubah
lewat import. Catatan verifikasi operator disimpan di file terpisah (verifikasi.csv)
sehingga perubahan catatan tidak pernah menyentuh hasil analisis.
"""
from __future__ import annotations

import io
import json
import numbers
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, time, timedelta
from pathlib import Path

import pandas as pd
import streamlit as st

from .config import DATA_DIR, STATUS_VERIFIKASI, TITIK

F_MONITORING = "monitoring.csv"
F_TEKANAN = "tekanan.csv"
F_EPISODE_INFO = "episode_info.csv"
F_KONTRIBUSI = "kontribusi.csv"
F_PROFIL = "profil_phase1.csv"
F_METADATA = "metadata_model.json"
F_VERIFIKASI = "verifikasi.csv"
F_VERIFIKASI_LOG = "verifikasi_log.csv"
F_RIWAYAT = "riwayat_import.json"
SEMUA_FILE = [F_MONITORING, F_TEKANAN, F_EPISODE_INFO, F_KONTRIBUSI, F_PROFIL,
              F_METADATA, F_VERIFIKASI, F_VERIFIKASI_LOG, F_RIWAYAT]

KOLOM_VERIFIKASI = ["Kunci_Episode", "Status", "Catatan", "Tanggal_Pemeriksaan", "Petugas", "Diperbarui"]


# --------------------------------------------------------------------------- util kolom

def _k(name) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


ALIAS = {
    "datetime": "Datetime", "waktu": "Datetime", "timestamp": "Datetime", "tanggaljam": "Datetime",
    "t2mewma": "T2_MEWMA", "t2": "T2_MEWMA", "t2statistic": "T2_MEWMA",
    "ucl": "UCL", "sinyal": "Sinyal", "signal": "Sinyal", "ooc": "Sinyal",
    "episode": "Episode", "episodeid": "Episode",
    "mulai": "Mulai", "selesai": "Selesai",
    "titikdominan": "Titik_Dominan", "dominan": "Titik_Dominan",
    "jumlahjamsinyal": "Jumlah_Jam_Sinyal", "jumlahjam": "Jumlah_Jam_Sinyal",
    "maksimumt2": "Maksimum_T2",
    "tanggal": "Tanggal", "jam": "Jam", "titik": "Titik",
    "ratarata": "Rata_rata", "mean": "Rata_rata", "rerata": "Rata_rata",
    "sd": "SD", "std": "SD", "n": "N", "jumlahdata": "N",
    **{_k(t): t for t in TITIK},
}


def normalisasi_kolom(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns={c: ALIAS.get(_k(c), c) for c in df.columns})


def _ke_angka(s: pd.Series) -> pd.Series:
    if s.dtype == object:
        s = s.astype(str).str.strip().str.replace(",", ".", regex=False)
    return pd.to_numeric(s, errors="coerce")


def _ke_bool(s: pd.Series) -> pd.Series:
    if s.dtype == bool:
        return s
    peta = {"true": True, "1": True, "ya": True, "yes": True, "false": False, "0": False, "tidak": False, "no": False}
    return s.astype(str).str.strip().str.lower().map(peta).fillna(False).astype(bool)


def _baca_tabel(raw: bytes, nama: str, header=0) -> pd.DataFrame:
    if nama.lower().endswith((".xlsx", ".xls")):
        return pd.read_excel(io.BytesIO(raw), header=header)
    teks = raw.decode("utf-8-sig", errors="replace")
    return pd.read_csv(io.StringIO(teks), sep=None, engine="python", header=header)


def _baca_tabel_tekanan(raw: bytes, nama: str) -> pd.DataFrame:
    """Baca tabel tekanan yang judulnya bisa lebih dari satu baris.

    Format data PDAM: baris 1 `Tanggal | Tekanan` (sel gabungan), baris 2 `Jam | IPA KM 8 | ... | Paldam`.
    Baris judul = baris pertama yang memuat nama titik; nama kolom yang kosong di baris itu diambil dari
    baris di atasnya (mis. `Tanggal` yang selnya digabung dua baris).
    """
    df = _baca_tabel(raw, nama, header=None).dropna(how="all").reset_index(drop=True)
    kunci_titik = {_k(t) for t in TITIK}
    judul = next((i for i in range(min(len(df), 10))
                  if sum(_k(v) in kunci_titik for v in df.iloc[i]) >= 3), 0)
    nama_kolom = []
    for j in range(df.shape[1]):
        atas = df.iloc[judul::-1, j].dropna()
        nama_kolom.append(str(atas.iloc[0]).strip() if not atas.empty else f"Kolom_{j + 1}")
    data = df.iloc[judul + 1:].reset_index(drop=True)
    data.columns = nama_kolom
    return data


def _jam_ke_timedelta(v) -> pd.Timedelta:
    """Jam dalam sehari dari berbagai bentuk: 01:00, 01.00, 1, 0.0417 (pecahan hari Excel), objek time."""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return pd.NaT
    if isinstance(v, (time, datetime)):
        return pd.Timedelta(hours=v.hour, minutes=v.minute, seconds=v.second)
    if isinstance(v, timedelta):
        return pd.Timedelta(v)
    if isinstance(v, numbers.Real):
        return pd.Timedelta(days=v) if 0 < v < 1 else pd.Timedelta(hours=v)
    s = str(v).strip()
    cocok = re.fullmatch(r"(\d{1,2})(?:[.:](\d{2}))?(?:[.:](\d{2}))?", s)
    if cocok:
        j, m, d = (int(x or 0) for x in cocok.groups())
        return pd.Timedelta(hours=j, minutes=m, seconds=d)
    try:
        return pd.Timedelta(pd.Timestamp(s).strftime("%H:%M:%S"))
    except (ValueError, TypeError):
        return pd.NaT


def _gabung_datetime(df: pd.DataFrame) -> pd.DataFrame:
    """Bentuk kolom Datetime dari Datetime atau pasangan Tanggal + Jam."""
    if "Datetime" not in df.columns and {"Tanggal", "Jam"}.issubset(df.columns):
        # Tanggal sering berupa sel gabungan (hanya terisi di jam pertama) → isi ke bawah.
        tanggal = pd.to_datetime(df["Tanggal"], errors="coerce", dayfirst=True).dt.normalize()
        tanggal = tanggal.where(df["Tanggal"].notna(), tanggal.ffill())
        df["Datetime"] = tanggal + df["Jam"].map(_jam_ke_timedelta).astype("timedelta64[ns]")
    if "Datetime" in df.columns:
        df["Datetime"] = pd.to_datetime(df["Datetime"], errors="coerce")
    return df


# --------------------------------------------------------------------------- baca / tulis

def _path(nama: str) -> Path:
    return DATA_DIR / nama


def _baca_csv(nama: str, tanggal=()) -> pd.DataFrame | None:
    p = _path(nama)
    if not p.exists():
        return None
    df = pd.read_csv(p)
    for c in tanggal:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors="coerce")
    return df


def _tulis_csv(nama: str, df: pd.DataFrame) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    df.to_csv(_path(nama), index=False)


def _baca_json(nama: str, default):
    p = _path(nama)
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else default


def _tulis_json(nama: str, obj) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    _path(nama).write_text(json.dumps(obj, indent=2, ensure_ascii=False, default=str), encoding="utf-8")


# --------------------------------------------------------------------------- episode

def hitung_episode(mon: pd.DataFrame) -> pd.Series:
    """Nomor episode = rangkaian sinyal berurutan per jam (aturan yang sama dengan skrip ekspor Colab)."""
    s = mon["Sinyal"].astype(bool)
    mulai = s & (~s.shift(fill_value=False) | mon["Datetime"].diff().gt(pd.Timedelta(hours=1)))
    return mulai.cumsum().where(s).astype("Int64")


def kunci_episode(mulai: pd.Timestamp) -> str:
    return pd.Timestamp(mulai).strftime("%Y-%m-%d %H:%M")


def _tabel_episode(mon: pd.DataFrame, info: pd.DataFrame | None, ver: pd.DataFrame) -> pd.DataFrame:
    sig = mon.loc[mon["Sinyal"]]
    if sig.empty:
        return pd.DataFrame(columns=["Episode", "Mulai", "Selesai", "Jumlah_Jam_Sinyal", "Maksimum_T2",
                                     "Waktu_T2_Maks", "Titik_Dominan", "Batch", "Kunci_Episode"] + KOLOM_VERIFIKASI[1:])
    ep = (sig.groupby("Episode")
          .agg(Mulai=("Datetime", "min"), Selesai=("Datetime", "max"),
               Jumlah_Jam_Sinyal=("Sinyal", "sum"), Maksimum_T2=("T2_MEWMA", "max"),
               Batch=("Batch", "first"))
          .reset_index())
    idx_maks = sig.groupby("Episode")["T2_MEWMA"].idxmax()
    ep["Waktu_T2_Maks"] = mon.loc[idx_maks.values, "Datetime"].values
    ep["Episode"] = ep["Episode"].astype(int)
    ep["Jumlah_Jam_Sinyal"] = ep["Jumlah_Jam_Sinyal"].astype(int)
    ep["Kunci_Episode"] = ep["Mulai"].map(kunci_episode)
    if info is not None and not info.empty:
        info = info.assign(Kunci_Episode=info["Mulai"].map(kunci_episode))[["Kunci_Episode", "Titik_Dominan"]]
        ep = ep.merge(info.drop_duplicates("Kunci_Episode", keep="last"), on="Kunci_Episode", how="left")
    else:
        ep["Titik_Dominan"] = pd.NA
    ep = ep.merge(ver, on="Kunci_Episode", how="left")
    ep["Status"] = ep["Status"].fillna(STATUS_VERIFIKASI[0])
    for c in ["Catatan", "Petugas", "Tanggal_Pemeriksaan", "Diperbarui"]:
        ep[c] = ep[c].fillna("")
    return ep


# --------------------------------------------------------------------------- store

@dataclass
class Store:
    metadata: dict
    monitoring: pd.DataFrame          # Datetime, T2_MEWMA, UCL, Sinyal, Batch, Episode
    tekanan: pd.DataFrame             # Datetime, 5 titik, Batch, Status_Model
    episodes: pd.DataFrame
    kontribusi: pd.DataFrame | None
    profil: pd.DataFrame | None       # profil Phase I yang diekspor (opsional)
    riwayat: list = field(default_factory=list)

    @property
    def kosong(self) -> bool:
        return self.monitoring.empty and self.tekanan.empty

    @property
    def ucl(self) -> float | None:
        if "ucl" in self.metadata:
            return float(self.metadata["ucl"])
        if not self.monitoring.empty:
            return float(self.monitoring["UCL"].iloc[0])
        return None


def _signature() -> tuple:
    return tuple((n, _path(n).stat().st_mtime_ns if _path(n).exists() else 0) for n in SEMUA_FILE)


@st.cache_data(show_spinner=False)
def _muat(_sig: tuple, data_dir: str) -> Store:
    meta = _baca_json(F_METADATA, {})
    mon = _baca_csv(F_MONITORING, ["Datetime"])
    if mon is None:
        mon = pd.DataFrame(columns=["Datetime", "T2_MEWMA", "UCL", "Sinyal", "Batch"])
    mon = mon.sort_values("Datetime").reset_index(drop=True)
    mon["Sinyal"] = _ke_bool(mon["Sinyal"]) if not mon.empty else mon["Sinyal"].astype(bool)
    mon["Episode"] = hitung_episode(mon) if not mon.empty else pd.Series(dtype="Int64")

    tek = _baca_csv(F_TEKANAN, ["Datetime"])
    if tek is None:
        tek = pd.DataFrame(columns=["Datetime", *TITIK, "Batch"])
    tek = tek.sort_values("Datetime").reset_index(drop=True)
    tek["Status_Model"] = "Belum diproses"
    tek.loc[tek["Datetime"].isin(set(mon["Datetime"])), "Status_Model"] = "Ada T²"
    fase1 = re.findall(r"\d{4}-\d{2}-\d{2}", str(meta.get("periode_phase_i", "")))
    if len(fase1) >= 2:
        di_fase1 = tek["Datetime"].between(pd.Timestamp(fase1[0]), pd.Timestamp(fase1[-1]) + pd.Timedelta(hours=23))
        tek.loc[di_fase1 & (tek["Status_Model"] != "Ada T²"), "Status_Model"] = "Referensi Phase I"

    info = _baca_csv(F_EPISODE_INFO, ["Mulai"])
    ver = _baca_csv(F_VERIFIKASI)
    if ver is None:
        ver = pd.DataFrame(columns=KOLOM_VERIFIKASI)
    ver = ver.astype({c: "object" for c in KOLOM_VERIFIKASI})
    eps = _tabel_episode(mon, info, ver)

    kon = _baca_csv(F_KONTRIBUSI, ["Datetime"])
    prof = _baca_csv(F_PROFIL)
    return Store(meta, mon, tek, eps, kon, prof, _baca_json(F_RIWAYAT, []))


def muat_store() -> Store:
    return _muat(_signature(), str(DATA_DIR))


def _segarkan():
    _muat.clear()


# --------------------------------------------------------------------------- import paket hasil ekspor

PETA_FILE = {
    "mewmamonitoringcsv": "monitoring", "monitoringcsv": "monitoring",
    "tekananaktualcsv": "tekanan",
    "episodemonitoringcsv": "episode", "episodescsv": "episode",
    "kontribusidiagnostikcsv": "kontribusi", "contributionscsv": "kontribusi",
    "phase1profilecsv": "profil", "profilphase1csv": "profil",
    "metadatamodeljson": "metadata", "modelmetadatajson": "metadata",
}


@dataclass
class Paket:
    monitoring: pd.DataFrame | None = None
    tekanan: pd.DataFrame | None = None
    episode: pd.DataFrame | None = None
    kontribusi: pd.DataFrame | None = None
    profil: pd.DataFrame | None = None
    metadata: dict | None = None
    file_dikenali: list = field(default_factory=list)
    galat: list = field(default_factory=list)       # menghalangi import
    peringatan: list = field(default_factory=list)  # boleh lanjut, tetapi perlu dibaca
    info: list = field(default_factory=list)


def _isi_file(unggahan) -> dict[str, bytes]:
    """Ubah daftar file unggahan (ZIP dan/atau file lepas) menjadi {nama_file: bytes}."""
    hasil = {}
    for f in unggahan:
        raw = f.getvalue()
        if f.name.lower().endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                for n in z.namelist():
                    if not n.endswith("/"):
                        hasil[Path(n).name] = z.read(n)
        else:
            hasil[f.name] = raw
    return hasil


def baca_paket(unggahan) -> Paket:
    pk = Paket()
    try:
        isi = _isi_file(unggahan)
    except zipfile.BadZipFile:
        pk.galat.append("File ZIP rusak atau bukan ZIP.")
        return pk

    for nama, raw in isi.items():
        jenis = PETA_FILE.get(_k(nama))
        if nama.upper().startswith("KONTRIBUSI_BELUM_TERSEDIA"):
            pk.info.append("Paket menyatakan data kontribusi belum diekspor dari notebook.")
            continue
        if jenis is None:
            pk.info.append(f"File `{nama}` tidak dikenali dan diabaikan.")
            continue
        pk.file_dikenali.append(nama)
        try:
            if jenis == "metadata":
                pk.metadata = json.loads(raw.decode("utf-8-sig"))
            elif jenis == "tekanan":
                pk.tekanan = normalisasi_kolom(_baca_tabel_tekanan(raw, nama))
            else:
                setattr(pk, jenis, normalisasi_kolom(_baca_tabel(raw, nama)))
        except Exception as e:  # noqa: BLE001 - laporkan apa pun yang gagal dibaca
            pk.galat.append(f"Gagal membaca `{nama}`: {e}")

    _rapikan_paket(pk)
    return pk


def _rapikan_paket(pk: Paket) -> None:
    if pk.monitoring is None:
        pk.galat.append("File monitoring (`mewma_monitoring.csv` atau `monitoring.csv`) wajib ada.")
        return
    m = _gabung_datetime(pk.monitoring)
    if not {"Datetime", "T2_MEWMA"}.issubset(m.columns):
        pk.galat.append("File monitoring harus memiliki kolom `Datetime` dan `T2_MEWMA`.")
        return
    # monitoring.csv versi brief memuat tekanan sekaligus; pisahkan bila tekanan_aktual.csv tidak ada.
    if pk.tekanan is None and set(TITIK).issubset(m.columns):
        pk.tekanan = m[["Datetime", *TITIK]].copy()

    m["T2_MEWMA"] = _ke_angka(m["T2_MEWMA"])
    ucl_meta = (pk.metadata or {}).get("ucl")
    if "UCL" not in m.columns:
        if ucl_meta is None:
            pk.galat.append("Nilai UCL tidak ditemukan (kolom `UCL` atau `ucl` di metadata).")
            return
        m["UCL"] = float(ucl_meta)
    m["UCL"] = _ke_angka(m["UCL"])
    m["Sinyal"] = _ke_bool(m["Sinyal"]) if "Sinyal" in m.columns else m["T2_MEWMA"] > m["UCL"]

    n_buruk = int(m["Datetime"].isna().sum() + m["T2_MEWMA"].isna().sum())
    if n_buruk:
        pk.peringatan.append(f"{n_buruk} baris monitoring memiliki waktu/T² kosong dan akan dibuang.")
    m = m.dropna(subset=["Datetime", "T2_MEWMA"]).sort_values("Datetime")
    if m["Datetime"].duplicated().any():
        pk.peringatan.append(f"{int(m['Datetime'].duplicated().sum())} waktu ganda di monitoring; yang terakhir dipakai.")
        m = m.drop_duplicates("Datetime", keep="last")
    if m["UCL"].nunique() > 1:
        pk.galat.append("Kolom UCL berisi lebih dari satu nilai; satu paket harus memakai satu UCL model.")
    if ucl_meta is not None and abs(m["UCL"].iloc[0] - float(ucl_meta)) > 1e-4:
        pk.galat.append(f"UCL di data ({m['UCL'].iloc[0]:.4f}) berbeda dengan metadata ({float(ucl_meta):.4f}).")
    beda_sinyal = int((m["Sinyal"] != (m["T2_MEWMA"] > m["UCL"])).sum())
    if beda_sinyal:
        pk.peringatan.append(f"{beda_sinyal} baris kolom `Sinyal` tidak sama dengan aturan T² > UCL. Data tetap dipakai apa adanya.")
    selang = m["Datetime"].diff().dropna()
    if (selang != pd.Timedelta(hours=1)).any():
        pk.peringatan.append(f"{int((selang != pd.Timedelta(hours=1)).sum())} selang waktu tidak sama dengan 1 jam (ada data hilang/loncat).")
    pk.monitoring = m[["Datetime", "T2_MEWMA", "UCL", "Sinyal"]].reset_index(drop=True)

    if pk.tekanan is not None:
        pk.tekanan, pesan = _rapikan_tekanan(pk.tekanan)
        pk.galat += pesan["galat"]
        pk.peringatan += pesan["peringatan"]

    if pk.episode is not None:
        e = pk.episode
        if "Mulai" in e.columns:
            e["Mulai"] = pd.to_datetime(e["Mulai"], errors="coerce")
            if "Titik_Dominan" not in e.columns:
                e["Titik_Dominan"] = pd.NA
            pk.episode = e[["Mulai", "Titik_Dominan"]].dropna(subset=["Mulai"])
        else:
            pk.peringatan.append("File episode tidak memiliki kolom `Mulai`; titik dominan tidak dapat dicocokkan.")
            pk.episode = None

    if pk.kontribusi is not None:
        k = _gabung_datetime(pk.kontribusi)
        if "Datetime" in k.columns and set(TITIK).issubset(k.columns):
            for t in TITIK:
                k[t] = _ke_angka(k[t])
            pk.kontribusi = k[["Datetime", *TITIK]].dropna(subset=["Datetime"])
        else:
            pk.peringatan.append("File kontribusi tidak memiliki `Datetime` + lima kolom titik; diabaikan.")
            pk.kontribusi = None

    if pk.profil is not None:
        p = pk.profil
        if {"Jam", "Titik", "Rata_rata", "SD"}.issubset(p.columns):
            pk.profil = p
        else:
            pk.peringatan.append("File profil Phase I harus berkolom Jam, Titik, Rata_rata, SD; diabaikan.")
            pk.profil = None


def _rapikan_tekanan(df: pd.DataFrame):
    pesan = {"galat": [], "peringatan": []}
    df = _gabung_datetime(normalisasi_kolom(df))
    kurang = [t for t in ["Datetime", *TITIK] if t not in df.columns]
    if kurang:
        pesan["galat"].append("Data tekanan kekurangan kolom: " + ", ".join(kurang)
                              + ". Nama titik harus: " + ", ".join(TITIK) + ".")
        return None, pesan
    for t in TITIK:
        df[t] = _ke_angka(df[t])
    n_waktu = int(df["Datetime"].isna().sum())
    if n_waktu:
        pesan["peringatan"].append(f"{n_waktu} baris tekanan tanpa waktu yang valid dibuang.")
    df = df.dropna(subset=["Datetime"]).sort_values("Datetime").drop_duplicates("Datetime", keep="last")
    kosong = df[TITIK].isna().sum()
    if kosong.any():
        pesan["peringatan"].append("Nilai tekanan kosong: " + ", ".join(f"{t} {n}" for t, n in kosong.items() if n))
    return df[["Datetime", *TITIK]].reset_index(drop=True), pesan


def cek_terhadap_store(pk: Paket, store: Store, mode: str) -> None:
    """Pemeriksaan kecocokan paket dengan data yang sudah tersimpan (mode 'tambah')."""
    if mode != "tambah" or store.monitoring.empty or pk.monitoring is None:
        return
    ucl_lama, ucl_baru = store.ucl, float(pk.monitoring["UCL"].iloc[0])
    if ucl_lama is not None and abs(ucl_lama - ucl_baru) > 1e-4:
        pk.galat.append(f"UCL paket ({ucl_baru:.4f}) berbeda dengan UCL data tersimpan ({ucl_lama:.4f}). "
                        "Hasil dari model berbeda tidak boleh digabung; gunakan mode 'Ganti seluruh data' bila memang model baru.")
    lam_lama, lam_baru = store.metadata.get("lambda_mewma"), (pk.metadata or {}).get("lambda_mewma")
    if lam_lama is not None and lam_baru is not None and float(lam_lama) != float(lam_baru):
        pk.galat.append(f"Lambda MEWMA paket ({lam_baru}) berbeda dengan data tersimpan ({lam_lama}).")
    tumpang = pk.monitoring["Datetime"].isin(set(store.monitoring["Datetime"])).sum()
    if tumpang:
        pk.peringatan.append(f"{int(tumpang)} waktu sudah ada di data tersimpan dan akan diganti dengan nilai dari paket ini.")
    akhir_lama = store.monitoring["Datetime"].max()
    awal_baru = pk.monitoring["Datetime"].min()
    if awal_baru > akhir_lama + pd.Timedelta(hours=1):
        jeda = (awal_baru - akhir_lama) / pd.Timedelta(hours=1)
        pk.info.append(f"Ada jeda {jeda:.0f} jam antara data tersimpan (s.d. {akhir_lama:%d %b %Y %H:%M}) "
                       f"dan paket baru (mulai {awal_baru:%d %b %Y %H:%M}).")
    elif awal_baru < store.monitoring["Datetime"].min():
        pk.peringatan.append("Paket berisi waktu yang lebih awal dari data tersimpan; nomor episode lama dapat bergeser "
                             "(catatan verifikasi tetap aman karena dikunci pada waktu mulai episode).")


def _gabung(lama: pd.DataFrame | None, baru: pd.DataFrame, kunci: str) -> pd.DataFrame:
    if lama is None or lama.empty:
        return baru.sort_values(kunci)
    return (pd.concat([lama, baru], ignore_index=True)
            .drop_duplicates(kunci, keep="last").sort_values(kunci))


def _catat_riwayat(entri: dict) -> None:
    r = _baca_json(F_RIWAYAT, [])
    r.append({"waktu_import": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), **entri})
    _tulis_json(F_RIWAYAT, r)


def simpan_paket(pk: Paket, mode: str, batch: str) -> None:
    """mode 'tambah' = gabung dengan data sebelumnya; 'ganti' = kosongkan hasil model lalu isi ulang."""
    if mode == "ganti":
        for n in [F_MONITORING, F_TEKANAN, F_EPISODE_INFO, F_KONTRIBUSI, F_PROFIL, F_METADATA]:
            _path(n).unlink(missing_ok=True)
    mon_lama = _baca_csv(F_MONITORING, ["Datetime"])
    _tulis_csv(F_MONITORING, _gabung(mon_lama, pk.monitoring.assign(Batch=batch), "Datetime"))
    if pk.tekanan is not None:
        _tulis_csv(F_TEKANAN, _gabung(_baca_csv(F_TEKANAN, ["Datetime"]), pk.tekanan.assign(Batch=batch), "Datetime"))
    if pk.episode is not None:
        _tulis_csv(F_EPISODE_INFO, _gabung(_baca_csv(F_EPISODE_INFO, ["Mulai"]), pk.episode, "Mulai"))
    if pk.kontribusi is not None:
        _tulis_csv(F_KONTRIBUSI, _gabung(_baca_csv(F_KONTRIBUSI, ["Datetime"]), pk.kontribusi, "Datetime"))
    if pk.profil is not None and (mode == "ganti" or not _path(F_PROFIL).exists()):
        _tulis_csv(F_PROFIL, pk.profil)
    if pk.metadata and (mode == "ganti" or not _path(F_METADATA).exists()):
        _tulis_json(F_METADATA, pk.metadata)
    m = pk.monitoring
    _catat_riwayat({"jenis": "Paket hasil model", "mode": mode, "batch": batch,
                    "file": ", ".join(pk.file_dikenali), "baris": len(m),
                    "periode": f"{m['Datetime'].min():%Y-%m-%d %H:%M} s.d. {m['Datetime'].max():%Y-%m-%d %H:%M}",
                    "sinyal": int(m["Sinyal"].sum())})
    _segarkan()


# --------------------------------------------------------------------------- import tekanan mentah

def baca_tekanan_mentah(unggahan):
    try:
        df = _baca_tabel_tekanan(unggahan.getvalue(), unggahan.name)
    except Exception as e:  # noqa: BLE001
        return None, {"galat": [f"Gagal membaca file: {e}"], "peringatan": []}
    return _rapikan_tekanan(df)


def simpan_tekanan_mentah(df: pd.DataFrame, batch: str, nama_file: str) -> None:
    _tulis_csv(F_TEKANAN, _gabung(_baca_csv(F_TEKANAN, ["Datetime"]), df.assign(Batch=batch), "Datetime"))
    _catat_riwayat({"jenis": "Tekanan mentah (tanpa T²)", "mode": "tambah", "batch": batch, "file": nama_file,
                    "baris": len(df),
                    "periode": f"{df['Datetime'].min():%Y-%m-%d %H:%M} s.d. {df['Datetime'].max():%Y-%m-%d %H:%M}",
                    "sinyal": None})
    _segarkan()


def template_tekanan() -> bytes:
    """Template Excel dengan tata letak data PDAM: judul dua baris (Tanggal | Tekanan → Jam + lima titik)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"], ws["B2"] = "Tanggal", "Tekanan", "Jam"
    for j, t in enumerate(TITIK, start=3):
        ws.cell(row=2, column=j, value=t)
    ws.merge_cells("A1:A2")
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=2 + len(TITIK))
    for baris in ws.iter_rows(min_row=1, max_row=2):
        for c in baris:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
    for jam in range(24):
        r = 3 + jam
        ws.cell(row=r, column=1, value=datetime(2026, 8, 1)).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=time(jam)).number_format = "hh:mm"
    ws.column_dimensions["A"].width = 13
    for j in range(3, 3 + len(TITIK)):
        ws.column_dimensions[ws.cell(row=2, column=j).column_letter].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- verifikasi

def simpan_verifikasi(kunci: str, status: str, catatan: str, tanggal: str, petugas: str) -> None:
    ver = _baca_csv(F_VERIFIKASI)
    if ver is None:
        ver = pd.DataFrame(columns=KOLOM_VERIFIKASI)
    baris = {"Kunci_Episode": kunci, "Status": status, "Catatan": catatan,
             "Tanggal_Pemeriksaan": tanggal, "Petugas": petugas,
             "Diperbarui": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    ver = pd.concat([ver[ver["Kunci_Episode"] != kunci], pd.DataFrame([baris])], ignore_index=True)
    _tulis_csv(F_VERIFIKASI, ver)
    log = _baca_csv(F_VERIFIKASI_LOG)
    _tulis_csv(F_VERIFIKASI_LOG, pd.concat([log, pd.DataFrame([baris])], ignore_index=True)
               if log is not None else pd.DataFrame([baris]))
    _segarkan()


def log_verifikasi() -> pd.DataFrame | None:
    return _baca_csv(F_VERIFIKASI_LOG)


# --------------------------------------------------------------------------- backup / pulihkan / reset

def buat_backup() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n in SEMUA_FILE:
            if _path(n).exists():
                z.write(_path(n), arcname=n)
    return buf.getvalue()


def pulihkan_backup(raw: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        nama = [n for n in z.namelist() if n in SEMUA_FILE]
        if F_MONITORING not in nama and F_TEKANAN not in nama:
            raise ValueError("ZIP ini bukan backup dashboard (tidak ada monitoring.csv / tekanan.csv). "
                             "Untuk paket hasil Colab, gunakan tab Import hasil model.")
        hapus_semua()
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        for n in nama:
            _path(n).write_bytes(z.read(n))
    _segarkan()
    return nama


def hapus_semua() -> None:
    if DATA_DIR.exists():
        for n in SEMUA_FILE:
            _path(n).unlink(missing_ok=True)
    _segarkan()


def ada_data() -> bool:
    return any(_path(n).exists() for n in [F_MONITORING, F_TEKANAN])

